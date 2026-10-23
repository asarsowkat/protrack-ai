#!/usr/bin/env python3
"""Representative tender costing workbooks for the Automatic WBS Classification tests (ProTrackAI v3.35).
Every figure is invented test data. Built with openpyxl; LibreOffice (headless) then recalculates and saves them, so the
files carry their formulas *and* their last calculated values, as Excel files do. File 13 is left as openpyxl writes it,
with formulas and no calculated values, to test 'formula result missing'.
Usage: python3 make_fixtures.py OUTDIR   (needs openpyxl and soffice)"""
import os, shutil, subprocess, sys, tempfile
import openpyxl
from openpyxl.styles import Font

OUT = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__)))
RAW = tempfile.mkdtemp()
B = Font(bold=True)

def book(): wb = openpyxl.Workbook(); return wb, wb.active
def head(ws, row, cols):
    for i, c in enumerate(cols, 1): ws.cell(row=row, column=i, value=c).font = B
def items(ws, start, rows, amt_col='F', q='C', r='E'):
    for i, x in enumerate(rows):
        n = start + i
        for j, v in enumerate(x, 1): ws.cell(row=n, column=j, value=v)
        ws[f'{amt_col}{n}'] = f'={q}{n}*{r}{n}'
    return start + len(rows)

# 1. a simple structured costing sheet
wb, ws = book(); ws.title = 'Costing'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
end = items(ws, 2, [[1, 'Supply of rebar Grade 60', 120, 'ton', 3200], [2, 'Rebar fixing crew', 4, 'month', 45000], [3, 'Excavation for cable trench', 850, 'cu.m', 38],
                    [4, 'Transformer supply 132/13.8kV 40MVA', 2, 'No.', 2850000], [5, 'Site villa rent', 12, 'month', 18000], [6, 'Recruitment fees', 1, 'LS', 60000]])
ws[f'B{end}'] = 'Total'; ws[f'F{end}'] = f'=SUM(F2:F{end-1})'
wb.save(f'{RAW}/01_simple.xlsx')

# 2. several differently structured sheets (header not in row 1, other column names)
wb, ws = book(); ws.title = 'Civil'
ws['A1'] = 'TENDER T-2026-118 — CIVIL WORKS'; ws['A2'] = 'Prepared by tender department (test data)'
head(ws, 4, ['S.No', 'Particulars', 'Quantity', 'UOM', 'Unit Rate', 'Total Price'])
end = items(ws, 5, [['1', 'Site preparation and clearing', 1, 'LS', 85000], ['2', 'Soil investigation boreholes', 6, 'nos', 9500], ['3', 'Backfill with selected material', 600, 'm3', 22],
                    ['4', 'Concrete C40 for foundations', 320, 'm3', 410]])
ws[f'B{end}'] = 'Total civil'; ws[f'F{end}'] = f'=SUM(F5:F{end-1})'
ws2 = wb.create_sheet('Electrical')
ws2['B2'] = 'Electrical'
head(ws2, 3, ['Ref', 'Scope', 'No. of units', 'Unit of measure', 'Price/Unit', 'Extended Price'])
end = items(ws2, 4, [['E1', 'XLPE cable 630mm2 132kV', 2400, 'm', 520], ['E2', 'Cable termination kits 132kV', 12, 'set', 8800], ['E3', 'Cable pulling technicians', 3, 'month', 52000]])
ws3 = wb.create_sheet('Site')
head(ws3, 1, ['Code', 'Cost description', 'Qty', 'Unit', 'Rate', 'Amount (SAR)'])
end = items(ws3, 2, [['SM-A', 'Project manager salary', 18, 'month', 42000], ['SM-B', 'Pick-up vehicles rental', 36, 'month', 3100], ['SM-C', 'Site office stationery and consumables', 18, 'month', 2500]])
wb.save(f'{RAW}/02_multi_sheet.xlsx')

# 3. hidden sheets and merged headings (a two-row header: "Price" over "Rate | Amount")
wb, ws = book(); ws.title = 'Summary'
ws.merge_cells('A1:F1'); ws['A1'] = 'BILL OF QUANTITIES — TESTING & COMMISSIONING'
ws['A2'] = 'Item'; ws['B2'] = 'Description'; ws['C2'] = 'Qty'; ws['D2'] = 'Unit'; ws.merge_cells('E2:F2'); ws['E2'] = 'Price'
ws['E3'] = 'Rate'; ws['F3'] = 'Amount'
ws.merge_cells('A4:F4'); ws['A4'] = 'A. TESTING AND COMMISSIONING'
end = items(ws, 5, [['A1', 'Commissioning of 132kV GIS bays', 6, 'No.', 65000], ['A2', 'Relay testing specialist services', 1, 'LS', 140000], ['A3', 'Testing equipment rental (primary injection set)', 3, 'month', 18000]])
ws[f'B{end}'] = 'Sub-total A'; ws[f'F{end}'] = f'=SUM(F5:F{end-1})'
h = wb.create_sheet('Calc'); h.sheet_state = 'hidden'
head(h, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(h, 2, [['C1', 'Engineering hours for protection design', 400, 'mh', 310]])
wb.save(f'{RAW}/03_hidden_merged.xlsx')

# 4. formula-driven subtotals and a grand total
wb, ws = book(); ws.title = 'BOQ'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
ws['A2'] = '1'; ws['B2'] = 'CIVIL WORKS'
items(ws, 3, [['1.1', 'Excavation in normal soil', 500, 'm3', 30], ['1.2', 'Backfilling and compaction', 420, 'm3', 25]])
ws['B5'] = 'Subtotal civil'; ws['F5'] = '=SUM(F3:F4)'
ws['A6'] = '2'; ws['B6'] = 'ELECTRICAL WORKS'
items(ws, 7, [['2.1', 'Power cable 300 sq mm 13.8kV', 1200, 'm', 160], ['2.2', 'LV switchgear panel', 2, 'No.', 95000]])
ws['B9'] = 'Subtotal electrical'; ws['F9'] = '=SUM(F7:F8)'
ws['B10'] = 'Grand Total'; ws['F10'] = '=F5+F9'
wb.save(f'{RAW}/04_formula_subtotals.xlsx')

# 5. a backup calculation repeated in a summary workbook
wb, ws = book(); ws.title = 'Summary'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
ws['A2'] = 1; ws['B2'] = 'Site staff salaries'; ws['C2'] = 1; ws['D2'] = 'LS'; ws['F2'] = "='Salaries backup'!F5"
ws['A3'] = 2; ws['B3'] = 'Insurance (CAR policy)'; ws['C3'] = 1; ws['D3'] = 'LS'; ws['E3'] = 75000; ws['F3'] = '=C3*E3'
ws['B4'] = 'Grand Total'; ws['F4'] = '=SUM(F2:F3)'
b = wb.create_sheet('Salaries backup')
head(b, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(b, 2, [['S1', 'Project manager', 18, 'month', 42000], ['S2', 'Site engineer', 36, 'month', 18000], ['S3', 'Safety officer', 18, 'month', 14000]])
b['B5'] = 'Total salaries'; b['F5'] = '=SUM(F2:F4)'
wb.save(f'{RAW}/05_backup_in_summary.xlsx')

# 6. similar descriptions with different specifications
wb, ws = book(); ws.title = 'Cables'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(ws, 2, [[1, '630 sq.mm XLPE cable 132kV', 1000, 'm', 520], [2, 'XLPE cable 630mm2 132kV', 1000, 'm', 520], [3, 'Power cable 630 sq mm 13.8kV', 1000, 'm', 210], [4, 'XLPE cable 300mm2 132kV', 800, 'm', 390]])
wb.save(f'{RAW}/06_similar_specs.xlsx')

# 7. combined material and labour costs (and a sheet with a material / labour breakdown)
wb, ws = book(); ws.title = 'Combined'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(ws, 2, [[1, 'Supply and install cable 132kV', 1, 'LS', 250000], [2, 'S&I earthing system', 1, 'LS', 48000], [3, 'Cable ladder, including installation', 300, 'm', 140]])
b = wb.create_sheet('Breakdown')
head(b, 1, ['Item', 'Description', 'Qty', 'Unit', 'Material amount', 'Labour amount', 'Total amount'])
b['A2'] = 1; b['B2'] = 'Supply and install cable 132kV'; b['C2'] = 1; b['D2'] = 'LS'; b['E2'] = 190000; b['F2'] = 60000; b['G2'] = '=E2+F2'
wb.save(f'{RAW}/07_combined.xlsx')

# 8. unknown cost descriptions
wb, ws = book(); ws.title = 'Misc'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(ws, 2, [[1, 'Bird diverters for overhead line', 120, 'No.', 260], [2, 'Miscellaneous allowance for local authority NOCs', 1, 'LS', 35000], [3, 'Zxq widget fee', 1, 'LS', 1200]])
wb.save(f'{RAW}/08_unknown.xlsx')

# 9. a WBS master import with invalid codes (blank code, duplicate code, unknown parent)
wb, ws = book(); ws.title = 'WBS master'
head(ws, 1, ['WBS code', 'Parent code', 'Description', 'Aliases', 'Keywords', 'Cost category', 'Project types', 'SAP WBS', 'Effective from', 'Effective to', 'Active', 'Postable', 'Mandatory review'])
for r in [['CIV-11', 'CIV', 'Dewatering', 'dewatering works', 'dewatering;pumping', 'Civil', '', '', '2026-01-01', '', 'Yes', 'Yes', 'No'],
          ['', 'CIV', 'No code given', '', '', 'Civil', '', '', '', '', 'Yes', 'Yes', 'No'],
          ['CIV-11', 'CIV', 'Dewatering again (duplicate code)', '', '', 'Civil', '', '', '', '', 'Yes', 'Yes', 'No'],
          ['XYZ-01', 'NOPE', 'Unknown parent', '', '', 'Other', '', '', '', '', 'Yes', 'Yes', 'No']]: ws.append(r)
wb.save(f'{RAW}/09_wbs_import_invalid.xlsx')

# 10. negative adjustments and a revision change (Rev 1 of file 1, with a credit)
wb, ws = book(); ws.title = 'Costing'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
end = items(ws, 2, [[1, 'Supply of rebar Grade 60', 130, 'ton', 3200], [2, 'Rebar fixing crew', 4, 'month', 45000], [3, 'Excavation for cable trench', 850, 'cu.m', 38],
                    [4, 'Transformer supply 132/13.8kV 40MVA', 2, 'No.', 2850000], [5, 'Site villa rent', 12, 'month', 18000], [6, 'Recruitment fees', 1, 'LS', 60000]])
ws[f'A{end}'] = 7; ws[f'B{end}'] = 'Discount on transformer supply'; ws[f'C{end}'] = 1; ws[f'D{end}'] = 'LS'; ws[f'E{end}'] = -150000; ws[f'F{end}'] = f'=C{end}*E{end}'
ws[f'B{end+1}'] = 'Total'; ws[f'F{end+1}'] = f'=SUM(F2:F{end})'
wb.save(f'{RAW}/10_negative_rev1.xlsx')

# 11. different currencies
wb, ws = book(); ws.title = 'Imported'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Currency', 'Amount'])
for i, (d, q, u, r, c) in enumerate([('GIS switchgear 132kV bay', 4, 'No.', 410000, 'EUR'), ('Protection relays', 24, 'No.', 6200, 'USD'), ('Local transport of switchgear', 1, 'LS', 38000, 'SAR')], 2):
    ws.append([i - 1, d, q, u, r, c]); ws[f'G{i}'] = f'=C{i}*E{i}'
wb.save(f'{RAW}/11_currencies.xlsx')

# 13. missing formula results: formulas saved without their calculated values (as written by some tools)
wb, ws = book(); ws.title = 'NoValues'
head(ws, 1, ['Item', 'Description', 'Qty', 'Unit', 'Rate', 'Amount'])
items(ws, 2, [[1, 'Concrete C40 for foundations', 100, 'm3', 410], [2, 'Formwork to foundations', 300, 'm2', 55]])
ws['B4'] = 'Total'; ws['F4'] = '=SUM(F2:F3)'
wb.save(f'{OUT}/13_missing_formula_results.xlsx')   # straight to OUT: not recalculated

# 16. a sheet without a recognisable header (needs a header mapping)
wb, ws = book(); ws.title = 'Rates'
ws.append(['Tender pricing (test data)']); ws.append([]); ws.append(['A', 'Cable laying manpower', 3, 'month', 41000, 123000]); ws.append(['B', 'Cable jointing manpower', 2, 'month', 47000, 94000])
wb.save(f'{RAW}/16_no_header.xlsx')


# 17. a WBS master import with order, SAP WBS segments and user fields (v3.36); no formulas, so written as is
wb, ws = book(); ws.title = 'WBS master'
head(ws, 1, ['Order', 'WBS code', 'Parent code', 'Description', 'SAP WBS segment', 'User field', 'Active', 'Postable'])
for r in [[5, 'SM2', 'SM', 'Vehicles', '002', 'SM', 'Yes', 'Yes'],
          [30, 'SM1-03', 'SM1', 'Salaries - Camp staff', '03', 'SM', 'Yes', 'Yes'],
          [40, 'SM1-04', 'SM1', 'Bad segment', 'A1', 'SM', 'Yes', 'Yes'],
          [50, 'SM1-05', 'SM1', 'Bad user field', '05', 'XYZ', 'Yes', 'Yes']]: ws.append(r)
wb.save(f'{OUT}/17_wbs_import_sap.xlsx')

# 18. the budget of an existing project, as kept for SAP (v3.37): actual SAP WBS IDs, user fields and three amount columns; no formulas
wb, ws = book(); ws.title = 'Budget'
ws['A1'] = 'PE-555 budget by SAP WBS element (test data)'
head(ws, 3, ['SAP WBS ID', 'WBS code', 'Description', 'User field', 'Cost type', 'Tender', 'Rev 0', 'Revised Rev 0'])
for r in [['PE-555-01-001-01', 'SM1-01', 'Salaries - Project office team', 'SM', 'Overhead', 1200000, 1250000, 1300000],
          ['PE-555-01-001-02', 'SM1-02', 'Salaries - Site office team', 'SM', 'Overhead', 800000, 820000, 840000],
          ['PE-555-04-001', 'ELM-01', 'Cables material', 'MTRL.ELC', 'Material', 5000000, 5100000, 4950000],
          ['PE-555-05-001', 'ELP-01', 'Cable laying manpower', 'MAN.ELC', 'Manpower', 900000, 910000, 905000.5],
          ['PE-555-13', 'PC', 'Prime cost items', 'PC ITEMS', None, 300000, 300000, 300000]]: ws.append(r)
ws.append([None, None, 'Total', None, None, 8200000, 8380000, 8295000.5])
wb.save(f'{OUT}/18_budget_import.xlsx')
# 19. a budget workbook with mistakes: a repeated SAP WBS ID, an amount that is not a number, an unknown user field, a row without an ID
wb, ws = book(); ws.title = 'Budget'
head(ws, 1, ['SAP WBS ID', 'Description', 'User field', 'Rev 0'])
for r in [['PE-555-01-001-01', 'Salaries - Project office team', 'SM', 1250000], ['PE-555-01-001-01', 'Salaries again', 'SM', 10],
          ['PE-555-04-001', 'Cables material', 'MTRL.ELC', 'TBC'], ['PE-555-09', 'Unknown field', 'XYZ', 100], [None, 'No ID here', 'SM', 5]]: ws.append(r)
wb.save(f'{OUT}/19_budget_import_invalid.xlsx')

# 20 to 24: monthly SAP ECC reports of project PE-555 (v3.38); no formulas, written as they are exported
def sap_wbs(path, period, rows, title='Budget/Actual/Commitment/Rem.Order Plan/Assigned'):
    wb, ws = book(); ws.title = 'S_ALR_87013558'
    ws['A1'] = title + ' (test data)'; ws['A2'] = f'Period: {period}  Currency: SAR'
    head(ws, 4, ['Object', 'Budget', 'Actual', 'Commitment', 'Assigned', 'Available'])
    for o, b, a, c in rows: ws.append([o, b, a, c, (a or 0) + (c or 0), (b or 0) - (a or 0) - (c or 0)])
    wb.save(path)
L08 = [('PE-555-01-001-01 Salaries - Project office team', 1250000, 600000, 100000), ('PE-555-01-001-02 Salaries - Site office team', 820000, 300000, 100000),
       ('PE-555-04-001 Cables material', 5100000, 3000000, 2500000), ('PE-555-05-001 Cable laying manpower', 900000, 820000, 0), ('PE-555-13 Prime cost items', 300000, 0, 0),
       ('PE-555-09-001 Earthing material', 0, 15000, 0)]
def tree(L):   # the report shows each level with the totals below it, as SAP does
    out = {}
    for o, b, a, c in L:
        i = o.split()[0]; out[i] = [o, b, a, c]
        p = i
        while p.count('-') > 1:
            p = p.rsplit('-', 1)[0]; t = out.setdefault(p, [p + {'PE-555-01': ' Site Management', 'PE-555-01-001': ' Salaries', 'PE-555-04': ' Electrical Material', 'PE-555-05': ' Electrical Manpower', 'PE-555-09': ' Earthing', 'PE-555': ' Project'}.get(p, ''), 0, 0, 0])
            if p != i: t[1] += b; t[2] += a; t[3] += c
    return [tuple(out[k]) for k in sorted(out, key=lambda k: [int(x) if x.isdigit() else x for x in k.split('-')])]
R08 = tree(L08)
sap_wbs(f'{OUT}/20_sap_wbs_2026-08.xlsx', '08/2026', R08 + [('Total', 8370000, 4735000, 2700000)])
sap_wbs(f'{OUT}/21_sap_wbs_2026-07.xlsx', '07/2026', tree([(o, b, round(a * 0.8), c) for o, b, a, c in L08]))
wb, ws = book(); ws.title = 'CJI3'
ws['A1'] = 'Display Actual Cost Line Items for Projects (test data)'
head(ws, 3, ['WBS Element', 'Cost Element', 'Cost element name', 'Posting Date', 'Ref Document Number', 'Val.in Obj.Crcy', 'Obj.Crcy'])
for r in [['PE-555-01-001-01', '610000', 'Salaries', '30.09.2026', '5100001', 50000, 'SAR'], ['PE-555-01-001-01', '610000', 'Salaries', '15.09.2026', '5100002', 25000, 'SAR'],
          ['PE-555-04-001', '410000', 'Materials', '20.09.2026', '5000123', 120000, 'SAR'], ['PE-555-04-001', '410000', 'Materials', '21.09.2026', '5000124', '1,500.00-', 'SAR'],
          ['PE-555-05-001', '620000', 'Subcontract labour', '28.09.2026', '5000200', 40000, 'SAR']]: ws.append(r)
ws.append([None, None, None, None, None, 233500, 'SAR'])
wb.save(f'{OUT}/22_cji3_2026-09.xlsx')
wb, ws = book(); ws.title = 'CJI5'
head(ws, 1, ['WBS Element', 'Purchasing Document', 'Cost Element', 'Value in Obj. Crcy', 'Currency'])
for r in [['PE-555-04-001', '4500011', '410000', 2380000, 'SAR'], ['PE-555-01-001-02', '4500020', '610000', 90000, 'SAR']]: ws.append(r)
wb.save(f'{OUT}/24_cji5_2026-09.xlsx')
wb, ws = book(); ws.title = 'Report'
head(ws, 1, ['WBS Element', 'Description', 'Budget', 'Actual', 'Currency'])
for r in [['PE-555-01-001-01', 'Salaries', 1250000, 'n/a', 'SAR'], ['PE-555-04-001', 'Cables', 5100000, 10, 'SAR'], ['PE-555-04-001', 'Cables again', 1, 1, 'SAR'], ['PE-555-05-001', 'Manpower', 900000, 5, 'USD'], [None, 'No WBS', 1, 1, 'SAR'], ['Total', None, 99, 99, None]]: ws.append(r)
wb.save(f'{OUT}/23_sap_invalid.xlsx')

# LibreOffice recalculates and saves with cached values; 01 also as .xls and 02 also as .xlsm (a macro-enabled container, no macros)
os.makedirs(OUT, exist_ok=True)
calc = tempfile.mkdtemp()
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', calc] + [f'{RAW}/{f}' for f in sorted(os.listdir(RAW))], check=True, capture_output=True)
for f in os.listdir(calc): shutil.copy(f'{calc}/{f}', f'{OUT}/{f}')
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xls', '--outdir', OUT, f'{calc}/01_simple.xlsx'], check=True, capture_output=True)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsm:Calc MS Excel 2007 VBA XML', '--outdir', OUT, f'{calc}/02_multi_sheet.xlsx'], check=True, capture_output=True)
with open(f'{OUT}/12_not_a_workbook.xlsx', 'w') as f: f.write('this is not a workbook')
print('\n'.join(sorted(os.listdir(OUT))))
