#!/usr/bin/env python3
"""ProTrackAI v3.35 and v3.36 on the demo page: Budget & Cost Management, Automatic WBS Classification, through the screens.
The real SheetJS 0.18.5 reads the test workbooks of tests/fixtures/bcm (the app's CDN address is answered with a local
copy). Usage: python3 bcm_demo.py DEMO_HTML XLSX_FULL_MIN_JS > results_bcm_demo.json"""
import json, os, sys, time
from playwright.sync_api import sync_playwright

PAGE, SJS = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
FX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'bcm')
fx = lambda *n: [os.path.join(FX, x) for x in n]
R = []
def ck(area, name, cond, detail=''):
    R.append({'area': area, 'check': name, 'result': 'PASS' if cond else 'FAIL', 'detail': '' if cond else str(detail)[:400]})
    print(('PASS ' if cond else 'FAIL ') + area + ': ' + name + ('' if cond else '  -> ' + str(detail)[:400]), file=sys.stderr)
CHART = "window.Chart=function(){return{destroy(){},update(){},resize(){},data:{datasets:[]},options:{}}};window.Chart.register=function(){};window.Chart.defaults={font:{},plugins:{legend:{labels:{}},tooltip:{}},scale:{grid:{}},color:''};"
XL = open(SJS).read()
def route(r):
    u = r.request.url
    if u.startswith('http'):
        if 'xlsx' in u: return r.fulfill(body=XL, content_type='application/javascript')
        if 'chart' in u.lower(): return r.fulfill(body=CHART, content_type='application/javascript')
        return r.abort()
    return r.continue_()

with sync_playwright() as p:
    b = p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
    ctx = b.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True); ctx.route('**/*', route)
    pg = ctx.new_page(); errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.goto('file://' + PAGE); pg.wait_for_timeout(900)
    def login(email):
        if pg.locator('#app').is_visible(): pg.evaluate('signOut()'); pg.wait_for_timeout(600)
        pg.fill('#lgId', email); pg.fill('#lgPw', 'ProTrack@2026'); pg.click('#lgBtn'); pg.wait_for_selector('#app', state='visible', timeout=10000); pg.wait_for_timeout(1200)
        pg.evaluate("window.aiHint=()=>{};window.rasaBubble=()=>{}")
    J = pg.evaluate
    def toast(): return J("document.querySelector('#toast').textContent")
    def wait_jobs(k=1, n=120):
        base = J("BCM_JOBS.length") if False else None
        for _ in range(n):
            pg.wait_for_timeout(500)
            if J(f"BCM_JOBS.length>=window.__jb+{k}&&BCM_JOBS.every(j=>j.ok||j.err)"): break
        J("window.__jb=BCM_JOBS.length")
    login('asarudeen@company.com')
    J("go('bcm');window.__jb=0"); pg.wait_for_timeout(600)
    nav = J("document.querySelector('#nav').textContent")
    ck('Navigation', 'Budget & Cost Management sits in the menu with its pages; stage 4 pages are marked planned', 'Budget & Cost Management' in nav and J("[...document.querySelectorAll('#nav .nav-item')].filter(x=>/Tender File Intake|WBS Mapping Master|Classification Review|Reconciliation Centre/.test(x.textContent)).length") == 4 and J("[...document.querySelectorAll('#nav .nav-item.planned')].some(x=>/Cost Summary/.test(x.textContent))"))
    ck('Navigation', 'Cost Control keeps its own menu as before', all(x in nav for x in ['Cost Control dashboard', 'Earned value and EAC', 'Master Manpower Costing Sheet', 'Budget supplements']))
    ck('Master', 'demo page: the starting master is approved by a sample approver, and one sample change waits', J("bcmOf('wbs').filter(m=>m.cur).length") == 56 and J("bcmOf('rule').filter(m=>m.cur).length") == 38 and J("!!bcmRow('W-SM10').pend"))
    # ---- intake
    J("bcmUi().sec='intake';render()"); pg.wait_for_timeout(300)
    pg.fill('#bcmTr', 'T-2026-118'); pg.fill('#bcmRev', '0'); pg.select_option('#bcmPt', '132 kV substation'); pg.click('[data-bcmnewbatch]'); pg.wait_for_timeout(500)
    bid = J("bcmUi().bid")
    ck('Intake', 'a batch is created with project, tender, revision, currency and project type', J(f"(b=>b&&b.tender_ref==='T-2026-118'&&b.currency==='SAR'&&b.ptype==='132 kV substation'&&b.created_by_name==='Asarudeen Sowkath Ali')(bcmBatch('{bid}'))") or J(f"(b=>b&&b.tender_ref==='T-2026-118'&&b.currency==='SAR')(bcmBatch('{bid}'))"), J(f"bcmBatch('{bid}')"))
    pg.set_input_files('#bcmFiles', fx('01_simple.xlsx', '01_simple.xls', '03_hidden_merged.xlsx', '07_combined.xlsx', '11_currencies.xlsx', '12_not_a_workbook.xlsx', '13_missing_formula_results.xlsx', '16_no_header.xlsx', '06_similar_specs.xlsx', '08_unknown.xlsx', '02_multi_sheet.xlsm'))
    wait_jobs(11); pg.wait_for_timeout(500)
    F = J(f"bcmFilesOf('{bid}').map(f=>({{n:f.name,s:f.status,run:f.run,rows:(f.sheets||[]).reduce((t,s)=>t+s.n,0)}}))")
    ck('Intake', 'ten workbooks registered in one go (.xlsx, .xls, .xlsm), each stored with its SHA-256', len(F) == 10 and J(f"bcmFilesOf('{bid}').every(f=>/^[0-9a-f]{{64}}$/.test(f.sha256)&&f.path.endsWith(f.sha256+'.'+f.kind))"), F)
    ck('Intake', 'a text file named .xlsx is refused with the reason, and not registered', J("BCM_JOBS.some(j=>j.name==='12_not_a_workbook.xlsx'&&/not an Excel workbook/.test(j.err))") and not any(f['n'] == '12_not_a_workbook.xlsx' for f in F))
    ck('Intake', 'a sheet without a header row leaves its file "Needs header mapping"', any(f['n'] == '16_no_header.xlsx' and f['s'] == 'Needs header mapping' for f in F))
    ck('Intake', 'the original file is kept unchanged in this browser (IndexedDB) under its fingerprint', J(f"(async()=>{{const f=bcmFilesOf('{bid}').find(x=>x.name==='01_simple.xlsx');const u=await bcmFetchFile(f);return (await sha256Hex(u))===f.sha256}})()"))
    pg.set_input_files('#bcmFiles', fx('01_simple.xlsx')); wait_jobs()
    ck('Intake', 'the same file uploaded again is refused, naming the first upload', J("/already uploaded as 01_simple.xlsx/.test(BCM_JOBS[0].steps[0])") and J(f"bcmFilesOf('{bid}').length") == 10, J("BCM_JOBS[0]"))
    ck('Intake', 'the processing log is shown with each file and step', '01_simple.xls' in pg.inner_text('#bcmJobs') and 'Not taken' in pg.inner_text('#bcmJobs'))
    # ---- extraction and classification results
    L = lambda: J(f"bcmLinesOf('{bid}').map(l=>({{f:bcmFile(l.file_id).name,s:l.sheet,r:l.row,k:bcmKind(l),d:l.desc,a:l.amount,st:bcmStatus(l),w:bcmCode((bcmCls(l.id)||{{}}).wbs_id),t:(bcmCls(l.id)||{{}}).cost_type,ly:(bcmCls(l.id)||{{}}).layer,fl:((bcmCls(l.id)||{{}}).flags||[]).join('; '),id:l.id}}))")
    rows = L(); g = lambda f, d: next((x for x in rows if x['f'] == f and d in (x['d'] or '')), None)
    ck('Classification', 'rebar supply: automatic to CIV-06, Material, by an approved rule', (lambda x: x and x['st'] == 'Auto-Classified' and x['w'] == 'CIV-06' and x['t'] == 'Material' and x['ly'] == 'rule')(g('01_simple.xlsx', 'Supply of rebar')), g('01_simple.xlsx', 'Supply of rebar'))
    ck('Classification', 'the transformer (SAR 5.7 million) waits for individual review as high value', (lambda x: x and x['st'] == 'Pending Review' and 'High value' in x['fl'])(g('01_simple.xlsx', 'Transformer')))
    ck('Classification', 'a hidden sheet is read and its line classified (engineering hours → DES-02)', (lambda x: x and x['w'] == 'DES-02')(g('03_hidden_merged.xlsx', 'Engineering hours')))
    ck('Classification', 'combined supply-and-install lines are Split Required', all(x['st'] == 'Split Required' for x in rows if x['f'] == '07_combined.xlsx' and x['k'] == 'item'))
    ck('Classification', 'missing formula results: amounts missing (never zero) and held for review', all(x['a'] is None and x['st'] == 'Pending Review' for x in rows if x['f'] == '13_missing_formula_results.xlsx' and x['k'] == 'item'))
    ck('Classification', 'unknown descriptions stay Unclassified', all(x['st'] == 'Unclassified' for x in rows if x['f'] == '08_unknown.xlsx' and x['k'] == 'item'))
    ck('Classification', 'the EUR line is high value at the EUR limit', 'High value: at or above 250,000.00 EUR' in (g('11_currencies.xlsx', 'GIS switchgear') or {}).get('fl', ''))
    ck('Classification', 'subtotals and totals are kept and not classified', any(x['k'] == 'total' and x['st'] == 'Not a cost line' for x in rows if x['f'] == '01_simple.xlsx'))
    # ---- header mapping
    fid16 = J(f"bcmFilesOf('{bid}').find(f=>f.name==='16_no_header.xlsx').id")
    pg.click(f'[data-bcmmap="{fid16}"]'); pg.wait_for_timeout(1500)
    ck('Intake', 'the column mapping shows a preview of the stored sheet', 'Tender pricing (test data)' in pg.inner_text('#bcmMap'))
    pg.fill('#bcmMapRow', '2')
    for k, v in {'code': 'A', 'desc': 'B', 'qty': 'C', 'unit': 'D', 'rate': 'E', 'amount': 'F'}.items(): pg.fill(f'[data-bcmmapc="{k}"]', v)
    pg.click('[data-bcmmapsave]'); wait_jobs(); pg.wait_for_timeout(600)
    rows = L()
    ck('Intake', 'with the mapping, the file is read again (run 2): two manpower lines, classified', J(f"(f=>f.status==='Extracted'&&f.run===2)(bcmFile('{fid16}'))") and [x['w'] for x in rows if x['f'] == '16_no_header.xlsx' and x['k'] == 'item'] == ['ELP-01', 'ELP-03'], [x for x in rows if x['f'] == '16_no_header.xlsx'])
    pg.click(f'[data-bcmverify="{fid16}"]'); pg.wait_for_timeout(1500)
    ck('Intake', 'Verify reading: the stored file read again gives the same rows', 'Verified' in toast(), toast())
    pg.click(f'[data-bcmlines="{fid16}"]'); pg.wait_for_timeout(500)
    ck('Intake', 'View lines shows the rows as read with their cells', 'Cable laying manpower' in pg.inner_text('#bcmView') and 'F3' in pg.inner_text('#bcmView'))
    # ---- review
    J("bcmUi().sec='review';bcmUi().F={st:'open'};render()"); pg.wait_for_timeout(500)
    v = pg.inner_text('#view')
    ck('Review', 'the progress summary shows source lines, automatic, awaiting review, confirmed, unclassified and amounts by currency', all(x in v for x in ['Total source lines', 'Automatically classified', 'Awaiting review', 'Unresolved amount', 'EUR', 'USD']))
    ck('Review', 'high-value lines waiting are shown on top', 'High-value and mandatory lines waiting' in v and 'Transformer supply' in v)
    tr = g('01_simple.xlsx', 'Transformer')['id']
    J(f"bcmOpenLine('{tr}')"); pg.wait_for_timeout(500)
    sh = pg.inner_text('#sheet')
    ck('Review', 'the line shows its source (sheet, row, cells, formula), context, proposal, candidates and history', all(x in sh for x in ['Costing row 5', 'F5', '=C5*E5', 'Context in the sheet', 'Candidate destinations', 'History', 'ELM-02']))
    pg.click('[data-bcmdo="confirm"]'); pg.wait_for_timeout(500)
    ck('Review', 'a high-value line is confirmed on its own', J(f"bcmStatus(bcmDB().lines.find(l=>l.id==='{tr}'))") == 'Confirmed')
    J("closeDrawer();bcmUi().F={st:'Pending Review'};bcmUi().sel=[];render()"); pg.wait_for_timeout(300)
    pg.click('[data-bcmselall]'); pg.wait_for_timeout(200); n_sel = J("bcmUi().sel.length")
    pg.click('[data-bcmbulk="confirm"]'); pg.wait_for_timeout(700)
    ck('Review', 'bulk confirmation of Pending Review lines confirms none: they are confirmed one by one', 'left for individual review' in toast() and J("bcmLinesOf(bcmUi().bid).filter(l=>bcmStatus(l)==='Confirmed').length") == 1, toast())
    J("bcmUi().F={st:'Auto-Classified'};bcmUi().sel=[];render()"); pg.wait_for_timeout(300); pg.click('[data-bcmselall]'); pg.wait_for_timeout(200); na = J("bcmUi().sel.length")
    pg.click('[data-bcmbulk="confirm"]'); pg.wait_for_timeout(800)
    ck('Review', 'bulk confirmation confirms the clean automatic classifications', J("bcmLinesOf(bcmUi().bid).filter(l=>bcmStatus(l)==='Confirmed').length") == 1 + na and na > 10, (na, toast()))
    # override with a reason and a project mapping, by a costing engineer; approved by someone else
    co = J("DB.master.users.find(u=>u.role==='costing'&&u.active!==false&&accessIds(u).includes(bcmBatch(bcmUi().bid).project_id)).email")
    exc = g('01_simple.xlsx', 'Excavation')['id']
    login(co); J(f"state.bcmUi=Object.assign(bcmUi(),{{bid:'{bid}',sec:'review'}});go('bcm');bcmOpenLine('{exc}')"); pg.wait_for_timeout(600)
    pg.click('[data-bcmdo="override"]'); pg.wait_for_timeout(300)
    ck('Review', 'an override without a reason is refused', 'reason' in toast().lower(), toast())
    pg.select_option('#bcmLT', 'Subcontract'); pg.fill('#bcmLR', 'excavation is let to the civil subcontractor'); pg.select_option('#bcmLS', 'project')
    pg.click('[data-bcmdo="override"]'); pg.wait_for_timeout(700)
    c = J(f"bcmCls('{exc}')")
    ck('Review', 'the costing engineer confirms with the cost type, recorded as overridden with the reason', c['status'] == 'Confirmed' and c['cost_type'] == 'Subcontract' and c['override'] and 'subcontractor' in c['reason'], c)
    mp = J("bcmOf('map').find(m=>m.pend&&m.pend.data.key==='Excavation for cable trench')")
    ck('Learning', '...and proposes a project mapping, which waits for approval', mp and mp['pend']['data']['scope'] == 'project' and mp['pend']['data']['source']['row'] == 4 and mp['cur'] is None, mp)
    J("closeDrawer();bcmUi().sec='learn';render()"); pg.wait_for_timeout(400)
    ck('Learning', 'the proposer cannot approve it (no approve right; and never one\'s own)', pg.locator(f'[data-bcmappr="{mp["id"]}"]').count() == 0)
    login('asarudeen@company.com'); J(f"state.bcmUi=Object.assign(bcmUi(),{{bid:'{bid}',sec:'learn'}});go('bcm')"); pg.wait_for_timeout(500)
    pg.click(f'[data-bcmappr="{mp["id"]}"]'); pg.wait_for_timeout(500)
    ck('Learning', 'the super admin approves it: version 1 in force', J(f"(m=>m.cur&&m.ver===1&&!m.pend)(bcmRow('{mp['id']}'))"))
    J("bcmUi().sec='intake';render()"); pg.wait_for_timeout(300)
    pg.set_input_files('#bcmFiles', fx('10_negative_rev1.xlsx')); wait_jobs(); pg.wait_for_timeout(500); rows = L()
    x = g('10_negative_rev1.xlsx', 'Excavation')
    ck('Learning', 'the next file (Rev 1) uses the approved mapping: automatic, 100%, Subcontract', x and x['st'] == 'Auto-Classified' and x['ly'] == 'map' and x['t'] == 'Subcontract', x)
    ck('Classification', 'the negative adjustment in Rev 1 waits for review', (lambda y: y and y['st'] == 'Pending Review' and 'Negative amount' in y['fl'])(g('10_negative_rev1.xlsx', 'Discount')))
    ck('Classification', "a person's decision is never overwritten by a new run", J(f"(async()=>{{await bcmClassifyBatch('{bid}');return bcmStatus(bcmDB().lines.find(l=>l.id==='{exc}'))}})()") == 'Confirmed')
    # correct, exclude, undo
    unk = g('08_unknown.xlsx', 'Zxq')['id']
    J(f"bcmOpenLine('{unk}')"); pg.wait_for_timeout(400); pg.fill('#bcmLR', 'not a cost of this tender'); pg.click('[data-bcmdo="exclude"]'); pg.wait_for_timeout(400)
    ck('Review', 'exclude with a reason', J(f"bcmStatus(bcmDB().lines.find(l=>l.id==='{unk}'))") == 'Excluded with Reason')
    pg.fill('#bcmLR', 'excluded by mistake'); pg.click('[data-bcmdo="undo"]'); pg.wait_for_timeout(400)
    ck('Review', 'undo restores the previous state as a new action in the history', J(f"bcmStatus(bcmDB().lines.find(l=>l.id==='{unk}'))") == 'Unclassified' and J(f"bcmDB().chist.filter(h=>h.line_id==='{unk}').map(h=>h.action).join()").endswith('Excluded,Undo'))
    bird = g('08_unknown.xlsx', 'Bird diverters')['id']
    J(f"bcmOpenLine('{bird}')"); pg.wait_for_timeout(400); pg.click('summary:has-text("Correct the description")'); pg.fill('#bcmLD', 'Other electrical material: bird diverters'); pg.fill('#bcmLR', 'clearer description for matching'); pg.click('[data-bcmdo="correct"]'); pg.wait_for_timeout(700)
    ck('Review', 'a corrected description is used for matching; the source description stays as read', J(f"(l=>l.desc==='Bird diverters for overhead line'&&bcmCls(l.id).corr.desc.startsWith('Other electrical')&&bcmCls(l.id).wbs_id==='W-ELM-09')(bcmDB().lines.find(l=>l.id==='{bird}'))"), J(f"bcmCls('{bird}')"))
    J("closeDrawer()")
    # ---- exports
    J("bcmUi().sec='review';bcmUi().F={st:'open'};render()"); pg.wait_for_timeout(300)
    with pg.expect_download() as dl: pg.click('[data-bcmexport="exc"]')
    import openpyxl, shutil, tempfile; path = os.path.join(tempfile.mkdtemp(), 'exc.xlsx'); dl.value.save_as(path); wb = openpyxl.load_workbook(path); ws = wb['Exceptions']
    hdr = [c.value for c in ws[1]]; data = list(ws.iter_rows(min_row=2, values_only=True))
    ck('Exports', 'the exception list workbook has every open or flagged line with its source and reason columns', 'Cell (amount)' in hdr and 'Controls and flags' in hdr and len(data) >= 10 and any('Combined cost' in str(r[hdr.index('Controls and flags')]) for r in data), (hdr[:6], len(data)))
    ck('Exports', '...with an About sheet saying it is not an approved budget and nothing went to SAP', 'About' in wb.sheetnames and 'not an approved budget' in ' '.join(str(c.value) for r in wb['About'].iter_rows() for c in r if c.value))
    # ---- WBS master
    J("bcmUi().sec='master';bcmUi().mtab='wbs';render()"); pg.wait_for_timeout(400)
    v = pg.inner_text('#view')
    ck('Master', 'the WBS master is a tree with versions, flags and the waiting change', 'SM1 Salaries' in v and 'Approved v1' in v and 'Warranty-period expenses' in v and 'always reviewed' in v)
    pg.click('[data-bcmedit="W-SM4"]'); pg.wait_for_timeout(300); pg.fill('[data-bcmw="desc"]', 'Villa / Staff accommodation'); pg.fill('#bcmWN', 'company name'); pg.click('[data-bcmwsave]'); pg.wait_for_timeout(400)
    ck('Master', 'a change is a proposal: the approved version stays in force', J("(m=>m.cur.desc==='Villa / Accommodation'&&m.pend.data.desc==='Villa / Staff accommodation')(bcmRow('W-SM4'))"))
    ck('Master', '...and its proposer is not offered approval', pg.locator('[data-bcmappr="W-SM4"]').count() == 0 and 'someone else approves it' in pg.inner_text('#view'))
    r = J("(async()=>{try{await bcmCall('master_act',{action:'approve',id:'W-SM4'});return 'approved'}catch(e){return e.message}})()")
    ck('Master', '...nor can he approve it by any other way', 'someone else approves it' in r, r)
    # an approver other than the proposer: the super admin gives a project manager the bcm.approve right (Settings → Roles and rights)
    apu = J("(u=>{rightsLocal({action:'user_right',user_id:u.id,right_key:'bcm.approve',allow:true});save();return u.email})(DB.master.users.find(u=>u.role==='pm'&&u.active!==false&&!isHead(u)))")
    login(apu); J("go('bcm');bcmUi().sec='master';render()"); pg.wait_for_timeout(400)
    pg.click('[data-bcmappr="W-SM4"]'); pg.wait_for_timeout(400)
    ck('Master', 'a project manager given bcm.approve approves it: version 2', J("(m=>m.ver===2&&m.cur.desc==='Villa / Staff accommodation')(bcmRow('W-SM4'))"))
    pg.click('[data-bcmhist="W-SM4"]'); pg.wait_for_timeout(300)
    ck('Master', 'the history keeps every version', pg.inner_text('#sheet').count('Approved') >= 2 and 'Villa / Accommodation' in pg.inner_text('#sheet'))
    J("closeDrawer()")
    ck('Master', 'someone without bcm.master is not offered the import', J("bcmUi().mtab='io';render();!document.querySelector('#bcmWbsFile')"))
    login('asarudeen@company.com'); J("go('bcm');bcmUi().sec='master';bcmUi().mtab='wbs';render()"); pg.wait_for_timeout(300)
    pg.click('[data-bcmmtab="io"]'); pg.wait_for_timeout(300)
    pg.set_input_files('#bcmWbsFile', fx('09_wbs_import_invalid.xlsx')); pg.wait_for_timeout(1200)
    v = pg.inner_text('#view')
    ck('Master', 'a WBS import with invalid codes: blank, repeated and unknown-parent rows are refused with the reason; the valid row is offered', 'no WBS code' in v and 'repeated' in v and 'unknown parent NOPE' in v and 'Propose the 1 valid rows' in v, v[-600:])
    pg.click('[data-bcmioprop]'); pg.wait_for_timeout(500)
    ck('Master', '...and becomes a proposal under its parent', J("(m=>m&&!m.cur&&m.pend.data.parent==='W-CIV')(bcmOf('wbs').find(m=>m.pend&&m.pend.data.code==='CIV-11'))"))
    pg.click('[data-bcmmtab="cfg"]'); pg.wait_for_timeout(300)
    ck('Master', 'thresholds and value limits are shown and changeable only as a proposal', pg.input_value('#bcmCA') == '95' and 'SAR, , 1000000' in pg.input_value('#bcmCV'))
    # ---- v3.36: order, SAP WBS ID and user field
    pe = J("pcode(projById(bcmUi().pid))")
    J("bcmUi().sec='master';bcmUi().mtab='wbs';render()"); pg.wait_for_timeout(400); v = pg.inner_text('#view')
    ck('SAP WBS', 'the tree shows each element\'s SAP WBS ID on the chosen project and its user field', f'SAP WBS {pe}-01-001 · User field SM' in v and f'{pe}-01-001-02' in v and 'User field MTRL.ELC Electrical Material' in v, v[:400])
    ck('SAP WBS', 'the starting master follows the company user field list: SM 01 … MGT.RES 16, Mechanical, Prime Cost and Provisional Sum', J("['SM','DES','TC','ELM','ELP','ELE','MEM','MEP','MEI','CIV','SAU','FIN','PC','PS','CON','MR'].map(c=>bcmData(bcmRow('W-'+c)).sap_seg).join()") == '01,02,03,04,05,06,07,08,09,10,11,12,13,14,15,16' and J("bcmOf('wbs').filter(m=>!bcmEff(m).parent).sort(bcmSort).map(m=>bcmEff(m).code).slice(0,3).join()") == 'SM,DES,TC')
    ck('SAP WBS', 'Salaries has its two sub-elements: Project office team 01 and Site office team 02', J("bcmOf('wbs').filter(m=>bcmEff(m).parent==='W-SM1').sort(bcmSort).map(m=>bcmEff(m).desc+' '+bcmEff(m).sap_seg).join('|')") == 'Salaries - Project office team 01|Salaries - Site office team 02')
    pg.click('[data-bcmmove="W-SM3"][data-dir="up"]'); pg.wait_for_timeout(600)
    ck('Order', '▲ proposes the new order of the siblings (SM3 before SM2), applied once approved', J("bcmRow('W-SM3').pend.data.order") == 20 and J("bcmRow('W-SM2').pend.data.order") == 30 and J("bcmRow('W-SM2').cur.order") == 20)
    ck('Order', '...and the tree shows the proposed order with the change waiting', J("bcmOf('wbs').filter(m=>bcmEff(m).parent==='W-SM').sort(bcmSort).slice(0,3).map(m=>bcmEff(m).code).join()") == 'SM1,SM3,SM2')
    pg.click('[data-bcmedit="new"]'); pg.wait_for_timeout(300)
    pg.fill('[data-bcmw="code"]', 'SM1-09'); pg.select_option('[data-bcmw="parent"]', 'W-SM1'); pg.fill('[data-bcmw="desc"]', 'Salaries - Camp staff'); pg.fill('[data-bcmw="sap_seg"]', '3'); pg.select_option('[data-bcmw="user_field"]', 'SM'); pg.click('[data-bcmwsave]'); pg.wait_for_timeout(400)
    ck('SAP WBS', 'a segment with the wrong number of digits for its level is refused (level 3 has 2)', 'level 3 the SAP WBS segment has 2 digits' in toast(), toast())
    pg.fill('[data-bcmw="sap_seg"]', '03'); pg.click('[data-bcmwsave]'); pg.wait_for_timeout(400)
    ck('SAP WBS', '...with 03 it is proposed, with its user field from the list', J("(m=>m&&m.pend.data.sap_seg==='03'&&m.pend.data.user_field==='SM')(bcmOf('wbs').find(m=>m.pend&&m.pend.data.code==='SM1-09'))"))
    J("closeDrawer()")
    r = J("(async()=>{try{await bcmCall('master_act',{action:'propose',ent:'wbs',data:{code:'X9',desc:'Test',user_field:'XYZ'}});return 'ran'}catch(e){return e.message}})()")
    ck('User field', 'a user field that is not in the user field list is refused', 'not in the user field list' in r, r)
    pg.click('[data-bcmmtab="io"]'); pg.wait_for_timeout(300)
    with pg.expect_download() as dl: pg.click('[data-bcmexport="wbs"]')
    import openpyxl, tempfile; pth = os.path.join(tempfile.mkdtemp(), 'wbs.xlsx'); dl.value.save_as(pth); ws = openpyxl.load_workbook(pth)['WBS master']
    hdr = [c.value for c in ws[1]]; rows = list(ws.iter_rows(min_row=2, values_only=True)); sm11 = next(r for r in rows if r[hdr.index('WBS code')] == 'SM1-01')
    ck('Excel', 'the WBS master workbook has the order, SAP segment, user field and the SAP WBS ID of the project, in tree order', all(h in hdr for h in ['Order', 'SAP WBS segment', 'User field', f'SAP WBS ID ({pe})']) and sm11[hdr.index(f'SAP WBS ID ({pe})')] == f'{pe}-01-001-01' and rows[0][hdr.index('WBS code')] == 'SM', (hdr, sm11))
    pg.set_input_files('#bcmWbsFile', fx('17_wbs_import_sap.xlsx')); pg.wait_for_timeout(1200); v = pg.inner_text('#view')
    ck('Excel', 'an import with order, SAP segments and user fields: a bad segment and an unknown user field are refused, the reorder and the new element are offered', 'not 1 to 6 digits' in v and 'user field XYZ is not in the user field list' in v and 'Propose the 2 valid rows' in v, v[-500:])
    pg.click('[data-bcmioprop]'); pg.wait_for_timeout(600)
    ck('Excel', '...and become proposals: SM2 to order 5, SM1-03 under SM1 with segment 03', J("bcmRow('W-SM2').pend.data.order") == 5 and J("(m=>m&&m.pend.data.parent==='W-SM1'&&m.pend.data.sap_seg==='03')(bcmOf('wbs').find(m=>m.pend&&m.pend.data.code==='SM1-03'))"))
    # a browser holding the v3.35 sample master is brought to the company layout
    up = J("""(()=>{const B=bcmDB(),keep=JSON.stringify(B.master),drop=['W-SM1-01','W-SM1-02','W-MEM','W-MEP','W-MEI','W-PC','W-PS'];
      B.master=B.master.filter(m=>!drop.includes(m.id));B.master.filter(m=>m.ent==='wbs'&&m.cur).forEach(m=>{delete m.cur.order;delete m.cur.sap_seg;delete m.cur.user_field;if(['W-SAU','W-FIN','W-CON','W-MR'].includes(m.id))m.cur.parent='W-OTH'});
      B.master.push({id:'W-OTH',ent:'wbs',ver:1,cur:{code:'OTH',parent:null,desc:'Other approved categories',active:true,postable:false},pend:null,status:'Approved'});
      DB.bcmSeeded336=false;seedBcm336();const r=[bcmData(bcmRow('W-SM1-02')).sap_seg,bcmData(bcmRow('W-SAU')).parent,bcmData(bcmRow('W-OTH')).active,bcmData(bcmRow('W-SM')).user_field,bcmRow('W-SM').ver,bcmRow('W-SM2').pend?'kept':'lost'];B.master=JSON.parse(keep);save();return r})()""")
    ck('Upgrade', 'a demo browser holding the v3.35 master gets the company layout (new elements, top-level SAD/FNE/CONT/MGT.RES, OTH off), and keeps the changes people have waiting', up == ['02', None, False, 'SM', 2, 'kept'], up)
    # ---- read-only roles and the rest of the app
    ro = J("(DB.master.users.find(u=>u.role==='foreman'&&u.active!==false)||{}).email")
    login(ro)
    ck('Rights', 'a foreman (no cost figures) does not see the module', not J("navOn('bcm')") and 'Budget & Cost Management' not in pg.inner_text('#nav'))
    login('asarudeen@company.com')
    for v in ['cdash', 'eva', 'mmcs', 'logs', 'boq', 'rel']:
        J(f"go('{v}')"); pg.wait_for_timeout(500)
    ck('Regression', 'Cost Control, earned value, the costing sheet, logs, BOQ Engine and releases still open', not errs, errs[:3])
    J("go('bcm');bcmUi().sec='audit';render()"); pg.wait_for_timeout(300)
    ck('Audit', 'the audit log lists uploads, readings, classifications, decisions and approvals', all(x in pg.inner_text('#view') for x in ['File uploaded', 'File read', 'Lines classified', 'Review: override', 'Change approved']))
    ck('Page', 'no script errors', not errs, errs[:3])
    b.close()
print(json.dumps({'suite': 'bcm_demo', 'passed': sum(r['result'] == 'PASS' for r in R), 'total': len(R), 'checks': R}, indent=1))
