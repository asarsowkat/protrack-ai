#!/usr/bin/env python3
"""ProTrackAI v3.38 on the demo page: Budget & Cost Management stage 4 (reporting and SAP preparation) through the screens:
monthly SAP reports (by WBS element, actual and commitment line items) imported, checked and frozen, column templates, the
Cost Summary with its 24 columns, the SAP loading workbook (6 sheets) of an approved version, SAP Loading Confirmed with
the SAP reference, a return, and a revision as supplements and returns. The real SheetJS 0.18.5 reads the workbooks.
Usage: python3 bcm_demo4.py DEMO_HTML XLSX_FULL_MIN_JS > results_bcm_demo4.json"""
import json, os, sys, tempfile
from playwright.sync_api import sync_playwright

PAGE, SJS = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
FX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'bcm')
fx = lambda *n: [os.path.join(FX, x) for x in n]
R = []
def ck(area, name, cond, detail=''):
    R.append({'area': area, 'check': name, 'result': 'PASS' if cond else 'FAIL', 'detail': '' if cond else str(detail)[:500]})
    print(('PASS ' if cond else 'FAIL ') + area + ': ' + name + ('' if cond else '  -> ' + str(detail)[:500]), file=sys.stderr)
CHART = "window.Chart=function(){return{destroy(){},update(){},resize(){},data:{datasets:[]},options:{}}};window.Chart.register=function(){};window.Chart.defaults={font:{},plugins:{legend:{labels:{}},tooltip:{}},scale:{grid:{}},color:''};"
XL = open(SJS).read()
def route(r):
    u = r.request.url
    if u.startswith('http'):
        if 'xlsx' in u: return r.fulfill(body=XL, content_type='application/javascript')
        if 'chart' in u.lower(): return r.fulfill(body=CHART, content_type='application/javascript')
        return r.abort()
    return r.continue_()

with sync_playwright() as p:
    b = p.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
    ctx = b.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True); ctx.route('**/*', route)
    pg = ctx.new_page(); errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
    pg.goto('file://' + PAGE); pg.wait_for_timeout(900)
    J = pg.evaluate
    def login(email):
        if pg.locator('#app').is_visible(): J('signOut()'); pg.wait_for_timeout(600)
        pg.fill('#lgId', email); pg.fill('#lgPw', 'ProTrack@2026'); pg.click('#lgBtn'); pg.wait_for_selector('#app', state='visible', timeout=10000); pg.wait_for_timeout(1200)
        J("window.aiHint=()=>{};window.rasaBubble=()=>{}")
    def toast(): return J("document.querySelector('#toast').textContent")
    def wait_jobs(k=1, n=120):
        for _ in range(n):
            pg.wait_for_timeout(500)
            if J(f"BCM_JOBS.length>=window.__jb+{k}&&BCM_JOBS.every(j=>j.ok||j.err)"): break
        J("window.__jb=BCM_JOBS.length")
    def sec(s, extra=''):
        J(f"closeDrawer();bcmUi().sec='{s}';{extra}render()"); pg.wait_for_timeout(400); return pg.inner_text('#view')
    import openpyxl
    def xl(dl):
        p_ = os.path.join(tempfile.mkdtemp(), 'x.xlsx'); dl.value.save_as(p_); return openpyxl.load_workbook(p_)
    def call3(fn, p): return J(f"(async()=>{{try{{return await bcmCall3('{fn}',{json.dumps(p)})}}catch(e){{return {{err:e.message}}}}}})()")
    login('asarudeen@company.com')
    J("go('bcm')"); pg.wait_for_timeout(600)
    nav = J("[...document.querySelectorAll('#nav .nav-item')].filter(x=>/Cost Summary|SAP Export/.test(x.textContent)).map(x=>x.textContent.trim().replace(/Planned.*$/,'')+(x.classList.contains('planned')?' [planned]':'')).join('|')")
    ck('Navigation', 'Cost Summary and SAP Export & Import are open; no page of the module is planned any more', 'Cost Summary' in nav and 'SAP Export & Import' in nav and '[planned]' not in nav and J("BCM_SECS.every(s=>!s[2])"), nav)
    pid = J("bcmUi().pid"); J(f"projById('{pid}').pe='PE-555';save()")
    co = J(f"DB.master.users.find(u=>u.role==='costing'&&u.active!==false&&accessIds(u).includes('{pid}')).email")
    rv = J(f"(u=>{{rightsLocal({{action:'user_right',user_id:u.id,right_key:'bcm.review',allow:true}});save();return u.email}})(DB.master.users.find(u=>u.role==='pm'&&u.active!==false&&accessIds(u).includes('{pid}')))")
    def budget(column, who_rev=rv):
        login(co); J(f"state.bcmUi=Object.assign(bcmUi(),{{pid:'{pid}',sec:'ver'}});go('bcm')"); pg.wait_for_timeout(400)
        pg.set_input_files('#bcmBudFile', fx('18_budget_import.xlsx')); pg.wait_for_timeout(1200); pg.select_option('[data-bcmbudcol]', label=column); pg.wait_for_timeout(300)
        pg.fill('#bcmBJ', column + ' as loaded'); pg.click('[data-bcmbudgo]'); pg.wait_for_timeout(700); J("closeDrawer()")
        vid = J(f"bcmVersOf('{pid}')[0].id"); call3('version', {'action': 'submit', 'id': vid})
        login(who_rev); call3('version', {'action': 'review', 'id': vid}); login('asarudeen@company.com'); call3('version', {'action': 'approve', 'id': vid})
        return vid
    V0 = budget('Rev 0')
    ck('Setup', 'Rev 0 of the project (PE-555) imported and approved in ProTrackAI', J(f"bcmS3().versions.find(v=>v.id==='{V0}').status") == 'Approved in ProTrackAI')
    # ---- SAP reports
    login(co); J(f"state.bcmUi=Object.assign(bcmUi(),{{pid:'{pid}',sec:'sap',sapTab:'imp'}});go('bcm')"); pg.wait_for_timeout(500)
    v = pg.inner_text('#view')
    ck('SAP reports', 'the page says there is no connection to SAP: reports come in as Excel, the workbook goes to the SAP team', 'There is no connection to SAP' in v and 'Import a monthly SAP report' in v)
    def imp(f, period, kind=None, extra=None):
        pg.fill('#bcmSapPer', period); pg.dispatch_event('#bcmSapPer', 'change'); pg.wait_for_timeout(200)
        pg.set_input_files('#bcmSapFile', fx(f)); pg.wait_for_timeout(600)
        try: pg.wait_for_selector('[data-bcmsapgo]', timeout=8000)
        except Exception: pass
        pg.wait_for_timeout(300)
        if kind: pg.select_option('#bcmSapKind', kind); pg.wait_for_timeout(300)
        if extra: extra()
        return pg.inner_text('#view')
    v = imp('21_sap_wbs_2026-07.xlsx', '2026-07')
    ck('SAP reports', 'a by-WBS report is read: header row 4, 12 rows, its totals', 'header in row 4' in v and '12 rows' in v and 'Budget SAR 8,370,000.00' in v and 'Actual SAR 3,788,000.00' in v and pg.locator('[data-bcmsapgo]').is_enabled(), v[-1500:])
    pg.click('[data-bcmsapgo]'); pg.wait_for_timeout(600)
    ck('SAP reports', 'imported and frozen', 'imported and frozen' in toast() and J(f"bcmRepsOf('{pid}').length") == 1, toast())
    v = imp('20_sap_wbs_2026-08.xlsx', '2026-08'); pg.click('summary:has-text("Columns")'); pg.fill('#bcmSapTplName', 'S_ALR_87013558 by WBS'); pg.click('[data-bcmsaptplsave]'); pg.wait_for_timeout(500)
    ck('Templates', 'the columns are saved as a template', J("bcmS4().sapTpls.map(t=>t.name+':'+t.map.wbs+','+t.map.budget).join()") == 'S_ALR_87013558 by WBS:Object,Budget', J("bcmS4().sapTpls"))
    pg.click('[data-bcmsapgo]'); pg.wait_for_timeout(600)
    v = imp('20_sap_wbs_2026-08.xlsx', '2026-08'); pg.click('[data-bcmsapgo]'); pg.wait_for_timeout(600)
    ck('SAP reports', 'the same file again is refused', 'already imported' in toast(), toast())
    J("bcmUi().sapImp=null;render()")
    v = imp('23_sap_invalid.xlsx', '2026-09')
    ck('SAP reports', 'a report with mistakes: each reason is shown and nothing can be imported', all(x in v for x in ['is not a number', 'in USD, not SAR', 'no WBS element', 'repeated in the report', 'the total row gives']) and pg.locator('[data-bcmsapgo]').is_disabled(), v[-1500:])
    v = imp('22_cji3_2026-09.xlsx', '2026-09')
    ck('SAP reports', 'actual cost line items (CJI3) are recognised, dates and the credit read', pg.input_value('#bcmSapKind') == 'actual' and '5 rows' in v and 'Val.in Obj.Crcy SAR 233,500.00' in v, v[-1200:])
    pg.click('[data-bcmsapgo]'); pg.wait_for_timeout(600)
    v = imp('24_cji5_2026-09.xlsx', '2026-09'); ck('SAP reports', 'open commitment line items (CJI5) are recognised', pg.input_value('#bcmSapKind') == 'commit')
    pg.click('[data-bcmsapgo]'); pg.wait_for_timeout(600)
    RR = J(f"bcmRepsOf('{pid}').map(r=>r.period+':'+r.kind+':'+r.status).sort().join('|')")
    ck('SAP reports', 'four reports frozen: July and August by WBS, September actual and commitment line items', RR == '2026-07:wbs:Frozen|2026-08:wbs:Frozen|2026-09:actual:Frozen|2026-09:commit:Frozen', RR)
    rid = J(f"bcmRepsOf('{pid}').find(r=>r.period==='2026-08').id"); J(f"bcmOpenSapRep('{rid}')"); pg.wait_for_timeout(500); sh = pg.inner_text('#sheet')
    ck('SAP reports', 'a report opens with its figures by WBS element (levels above showing only their own part) and its rows as imported', 'its own part only' in sh and 'Rows as imported (12)' in sh, sh[:400])
    J("closeDrawer()")
    # ---- Cost Summary
    J("bcmUi().sec='sum';render()"); pg.wait_for_timeout(900); v = pg.inner_text('#view')
    ck('Cost summary', 'the latest period (2026-09) is shown with the 24 columns', pg.input_value('[data-bcmcsper]') == '2026-09' and J("document.querySelectorAll('table.bcmcs thead th').length") == 24, pg.input_value('[data-bcmcsper]'))
    T = J("(t=>({c:BCME.decStr(t.current),a:BCME.decStr(t.actual),m:BCME.decStr(t.commit),e:BCME.decStr(t.eac),v:BCME.decStr(t.vac)}))(bcmUi().csLast.total)")
    ck('Cost summary', 'the project: budget 8,380,000, actual 4,968,500, commitments 2,470,000, EAC 8,793,500, VAC -413,500', T == {'c': '8380000', 'a': '4968500', 'm': '2470000', 'e': '8793500', 'v': '-413500'}, T)
    ck('Cost summary', 'the sources are said (August by WBS + September line items; commitments from the line items)', 'Budget, actual and commitment by WBS element 2026-08 + Actual cost line items 2026-09' in v and 'Commitments: Open commitment line items 2026-09' in v, v[:1500])
    st = J("Object.fromEntries(bcmUi().csLast.rows.filter(r=>!r.kids).map(r=>[r.sap,r.c.status.join(',')]))")
    ck('Cost summary', 'over budget, at risk with the SAP budget differing, and cost outside the approved budget are flagged', st.get('PE-555-04-001') == 'Over budget' and st.get('PE-555-05-001') == 'At risk,SAP budget differs' and st.get('PE-555-09-001') == 'Not in the approved budget', st)
    pg.click('[data-bcmcsf="Over budget"]'); pg.wait_for_timeout(300)
    ck('Cost summary', 'the filter shows the elements over budget and their levels', J("[...document.querySelectorAll('table.bcmcs tbody tr')].map(r=>r.cells[0].textContent.trim()).slice(0,-1).join()") == 'PE-555-04,PE-555-04-001')
    pg.click('[data-bcmcsf="all"]'); pg.wait_for_timeout(200)
    with pg.expect_download() as dl: pg.click('[data-bcmcsexp]')
    wb = xl(dl); ws = wb['Cost summary']; hdr = [c.value for c in ws[1]]
    ck('Cost summary', 'the Excel export has the 24 columns, the definitions and the sources', len(hdr) == 24 and hdr[20] == 'Estimate at completion' and 'Definitions' in wb.sheetnames and 'Sources' in wb.sheetnames, hdr)
    pg.select_option('[data-bcmcsper]', '2026-08'); pg.wait_for_timeout(900)
    ap = J("BCME.decStr(bcmUi().csLast.rows.find(r=>r.sap==='PE-555-01-001-01').c.actual_p)")
    ck('Cost summary', 'August: the actual of the month from the difference of the July and August reports', ap == '120000' and 'difference between the by-WBS reports' in pg.inner_text('#view'), ap)
    # ---- SAP loading
    J("bcmUi().sec='sap';bcmUi().sapTab='load';render()"); pg.wait_for_timeout(400)
    ck('SAP loading', 'a costing engineer without bcm.sap is not offered the workbook', pg.locator(f'[data-bcmsapgen="{V0}"]').count() == 0 and 'you do not have that right' in pg.inner_text('#view'))
    r = J(f"(async()=>{{try{{await bcmCall4('sap_export',{{action:'generate',version_id:'{V0}'}});return 'made'}}catch(e){{return e.message}}}})()")
    ck('SAP loading', '...nor can he make it any other way', 'bcm.sap' in r, r)
    login('asarudeen@company.com'); J(f"state.bcmUi=Object.assign(bcmUi(),{{pid:'{pid}',sec:'sap',sapTab:'load'}});go('bcm')"); pg.wait_for_timeout(500)
    with pg.expect_download() as dl: pg.click(f'[data-bcmsapgen="{V0}"]')
    wb = xl(dl)
    ck('SAP loading', 'the SAP authorized user prepares the workbook: six sheets, the original budget by SAP WBS ID', wb.sheetnames == ['Header', 'WBS Elements', 'Original Budget', 'Supplements & Returns', 'Control Totals', 'Approvals & Checks']
       and [r[0] for r in wb['Original Budget'].iter_rows(min_row=2, values_only=True)] == ['PE-555-01-001-01', 'PE-555-01-001-02', 'PE-555-04-001', 'PE-555-05-001', 'PE-555-13'] and wb['Original Budget']['C5'].value == 910000, wb.sheetnames)
    ck('SAP loading', '...with the levels above in WBS Elements, the control totals and the approvals', any(r[0] == 'PE-555-01-001' and r[2] == 'PE-555-01' for r in wb['WBS Elements'].iter_rows(values_only=True)) and any(r[0] == 'Total of this version' and r[2] == 8380000 for r in wb['Control Totals'].iter_rows(values_only=True))
       and any('Approve' in str(r[1]) for r in wb['Approvals & Checks'].iter_rows(values_only=True)))
    ck('SAP loading', 'the version is Ready for SAP Loading', J(f"bcmS3().versions.find(v=>v.id==='{V0}').status") == 'Ready for SAP Loading')
    eid = J(f"bcmExpsOf('{V0}').slice(-1)[0].id")
    pg.fill(f'[data-bcmsapref="{eid}"]', 'x'); pg.click(f'[data-bcmsapok="{eid}"]'); pg.wait_for_timeout(400)
    ck('SAP loading', 'a loading without the SAP reference is refused', 'SAP reference' in toast(), toast())
    pg.fill(f'[data-bcmsapref="{eid}"]', 'CJ30-4711'); pg.click(f'[data-bcmsapok="{eid}"]'); pg.wait_for_timeout(500)
    ck('SAP loading', 'SAP Loading Confirmed is recorded with the SAP reference', J(f"(v=>v.status+'/'+v.sap_ref)(bcmS3().versions.find(v=>v.id==='{V0}'))") == 'SAP Loading Confirmed/CJ30-4711' and 'Last loaded in SAP: Rev 0 0 (SAP reference CJ30-4711)' in pg.inner_text('#view'))
    V1 = budget('Revised Rev 0')
    login('asarudeen@company.com'); J(f"state.bcmUi=Object.assign(bcmUi(),{{pid:'{pid}',sec:'sap',sapTab:'load'}});go('bcm')"); pg.wait_for_timeout(500)
    with pg.expect_download() as dl: pg.click(f'[data-bcmsapgen="{V1}"]')
    wb = xl(dl); sr = list(wb['Supplements & Returns'].iter_rows(min_row=2, values_only=True))
    ck('SAP loading', 'the revision goes as supplements and returns against Rev 0 (CJ30-4711)', [(r[0], r[4], r[5]) for r in sr] == [('PE-555-01-001-01', 50000, 'Supplement'), ('PE-555-01-001-02', 20000, 'Supplement'), ('PE-555-04-001', -150000, 'Return'), ('PE-555-05-001', -4999.5, 'Return')]
       and 'Not used' in str(wb['Original Budget']['A1'].value) and any(r[2] == 'OK' for r in wb['Control Totals'].iter_rows(values_only=True)), sr)
    e1 = J(f"bcmExpsOf('{V1}').slice(-1)[0].id")
    pg.fill(f'[data-bcmsaprs="{e1}"]', 'fiscal year closed in SAP'); pg.click(f'[data-bcmsapret="{e1}"]'); pg.wait_for_timeout(500)
    ck('SAP loading', 'returned with the reason: the version is approved in ProTrackAI again', J(f"bcmS3().versions.find(v=>v.id==='{V1}').status") == 'Approved in ProTrackAI' and J(f"bcmExpsOf('{V1}')[0].status") == 'Returned')
    with pg.expect_download() as dl: pg.click(f'[data-bcmsapgen="{V1}"]')
    e2 = J(f"bcmExpsOf('{V1}').slice(-1)[0]")
    pg.fill(f'[data-bcmsapref="{e2["id"]}"]', 'CJ37-0815'); pg.click(f'[data-bcmsapok="{e2["id"]}"]'); pg.wait_for_timeout(500)
    ck('SAP loading', 'workbook 2 confirmed: the revision is loaded (CJ37-0815)', e2['seq'] == 2 and J(f"bcmS3().versions.find(v=>v.id==='{V1}').sap_ref") == 'CJ37-0815')
    v = pg.inner_text('#view')
    ck('SAP loading', 'the list of workbooks keeps all three with their status and reference', all(x in v for x in ['CJ30-4711', 'CJ37-0815', 'Returned', 'fiscal year closed']))
    # ---- after the revision
    J("bcmUi().sec='sum';bcmUi().csPer='2026-09';render()"); pg.wait_for_timeout(900)
    c = J("(t=>BCME.decStr(t.current))(bcmUi().csLast.total)"); st = J("Object.fromEntries(bcmUi().csLast.rows.filter(r=>!r.kids).map(r=>[r.sap,r.c.status.join(',')]))")
    ck('Cost summary', 'the current budget is now the Revised Rev 0 (8,295,000.50); the SAP budget of the reports (Rev 0) differs from it', c == '8295000.5' and 'SAP budget differs' in st.get('PE-555-01-001-01', ''), (c, st))
    v = sec('appr')
    ck('Approval history', 'lists the SAP loadings and the return', all(x in v for x in ['SAP loading confirmed', 'SAP loading workbook returned', 'SAP loading workbook prepared']))
    v = sec('ov')
    ck('Overview', 'the overview has the SAP loading tile instead of a planned one', 'Ready for SAP loading' in v and J("[...document.querySelectorAll('.boqkpi')].filter(k=>/Stage/.test(k.textContent)).length") == 0)
    v = sec('audit')
    ck('Audit', 'the audit log has the imports, the template and the SAP loading', all(x in v for x in ['SAP report imported', 'SAP report template saved', 'SAP loading confirmed']))
    ck('Page', 'no script errors', not errs, errs[:3])
    b.close()
print(json.dumps({'suite': 'bcm_demo4', 'passed': sum(r['result'] == 'PASS' for r in R), 'total': len(R), 'checks': R}, indent=1))
